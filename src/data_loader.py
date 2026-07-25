#data_loader.py
#
# Excel / Google スプレッドシートから作業リストを DataFrame として読み込む

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from task_image import pt_to_px

DEFAULT_COL_WIDTH = 120


def convert_deadline_value(x):
    if pd.isna(x):
        return pd.NaT
    if isinstance(x, (int, float)):
        try:
            return (pd.to_datetime("1899-12-30") + pd.to_timedelta(x, unit="D"))
        except Exception:
            return pd.to_datetime(x, errors="coerce")
    try:
        return pd.to_datetime(x, errors="coerce")
    except Exception:
        return pd.NaT


def load_dataframe_from_excel(excel_path: str):
    """
    Excelファイル(「作業リスト」シート、C～K列)を読み込み DataFrame を返す。
    戻り値: (df, col_width_map, row_height_base)
    """
    df = pd.read_excel(
        excel_path,
        sheet_name="作業リスト",
        usecols="C:K"
    )

    wb = load_workbook(excel_path)
    ws = wb["作業リスト"]

    col_width_map = {}
    start_col_index = 3

    for i, col_name in enumerate(df.columns):
        excel_col_index = start_col_index + i
        letter = get_column_letter(excel_col_index)
        dim = ws.column_dimensions.get(letter)
        if dim and dim.width:
            col_width_map[col_name] = int(dim.width * 8.2 + 12)
        else:
            col_width_map[col_name] = DEFAULT_COL_WIDTH

    data_row_index = 8
    excel_row_height = ws.row_dimensions[data_row_index].height
    if excel_row_height:
        row_height_base = int(excel_row_height * 96 / 72) + 3
    else:
        row_height_base = pt_to_px(15)

    return df, col_width_map, row_height_base


def load_dataframe_from_sheets(spreadsheet_url: str, sheet_name: str = "作業リスト"):
    """
    Google Sheets APIでスプレッドシートを読み込み DataFrameを返す。
    トークンは google_auth_helper から自動取得（事前に GUI で認可が必要）。
    戻り値: (df, col_width_map, row_height_base)
    """
    try:
        import gspread
        from google_auth_helper import get_creds, has_token
    except ImportError as e:
        raise ImportError(
            "gspread または google-auth がインストールされていません。\n"
            "pip install gspread google-auth google-auth-oauthlib を実行してください。\n" + str(e)
        )

    if not has_token():
        raise ValueError(
            "Google Sheets の認可がまだ完了していません。\n"
            "GUI の「Google で認可する」ボタンをクリックして認可してください。"
        )

    try:
        creds = get_creds()
        if not creds:
            raise ValueError("トークンの取得に失敗しました。もう一度認可してください。")
        gc = gspread.authorize(creds)
    except Exception as e:
        raise ValueError(f"Google Sheets の認可に失敗しました:\n{str(e)}")

    # URLからスプレッドシートを開く
    sh = gc.open_by_url(spreadsheet_url)
    ws = sh.worksheet(sheet_name)

    all_values = ws.get_all_values()

    if not all_values:
        raise ValueError("スプレッドシートにデータがありません")

    # C列（index=2）から K列（index=10）に相当する列を取得
    # ヘッダー行を特定（空でない最初の行）
    header_row_idx = 0
    for i, row in enumerate(all_values):
        if any(cell.strip() for cell in row[2:11]):
            header_row_idx = i
            break

    headers = all_values[header_row_idx][2:11]
    data_rows = [r[2:11] for r in all_values[header_row_idx + 1:] if any(c.strip() for c in r[2:11])]

    df = pd.DataFrame(data_rows, columns=headers)

    col_width_map = {h: DEFAULT_COL_WIDTH for h in headers}
    for col, w in {"内容": 160, "詳細": 200, "備考": 140, "担当": 80, "締切": 70}.items():
        if col in col_width_map:
            col_width_map[col] = w

    row_height_base = 24

    return df, col_width_map, row_height_base
